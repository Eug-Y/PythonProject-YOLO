"""Create the final Data_v8 YOLO11 technical decision brief as a DOCX."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SUMMARY = ROOT / "reports" / "Data_v8_yolo11_final_report" / "summary.json"
DEPLOYMENT_IMAGE = ROOT / "reports" / "Data_v8_yolo11_final_report" / "assets" / "deployment_overview.jpg"
C1_IMAGE = ROOT / "reports" / "Data_v8_yolo11_final_report" / "assets" / "c1_gate_overview.jpg"
OUTPUT = ROOT / "reports" / "项目最终进度与模型效果总结_Data_v8.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "65736C"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "FFF4D9"
GREEN = "237A4D"
RED = "9B1C1C"


def set_run_font(run, *, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("17231E")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != 9360:
        raise ValueError(f"Table widths must sum to 9360 DXA: {widths}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = 0
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        shade_cell(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, size=9.2, bold=True, color=INK)
    for values in rows:
        cells = table.add_row().cells
        for index, text in enumerate(values):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            if index > 0 and len(text) < 22:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            set_run_font(run, size=9.1)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(4)
    return table


def create_numbering(doc: Document, *, ordered: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    level.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1." if ordered else "•")
    level.append(text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    level.append(ppr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc: Document, text: str, num_id: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    ppr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(num)
    ppr.append(num_pr)
    run = p.add_run(text)
    set_run_font(run, size=11)


def add_callout(doc: Document, label: str, text: str, *, fill: str = CALLOUT, color: str = INK) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.line_spacing = 1.10
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "A87800" if fill == CALLOUT else color)
    borders.append(left)
    ppr.append(borders)
    lead = p.add_run(f"{label}：")
    set_run_font(lead, size=11, bold=True, color=color)
    body = p.add_run(text)
    set_run_font(body, size=11, color=color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=8.5, color=MUTED)


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = p.add_run("DATA_V8 FINAL MODEL REVIEW")
    set_run_font(left, size=8, bold=True, color=MUTED)
    right = p.add_run("\tVALID-ONLY / 2026-08-17")
    set_run_font(right, size=8, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_picture(doc: Document, path: Path, caption: str, alt: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(6.5))
    inline._inline.docPr.set("descr", alt)
    cap = doc.add_paragraph(caption, style="Caption")
    cap.paragraph_format.keep_with_next = False


def add_body_paragraph(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = p.add_run(text[len(bold_prefix) :])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)


def main() -> int:
    for path in (SUMMARY, DEPLOYMENT_IMAGE, C1_IMAGE):
        if not path.is_file():
            raise FileNotFoundError(path)
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    if summary.get("held_out_test_used") is not False or summary.get("gate_passed") is not False:
        raise RuntimeError("Final Word brief requires the completed valid-only failed gate")

    doc = Document()
    style_document(doc)
    add_header_footer(doc)
    doc.core_properties.title = "项目最终进度与模型效果总结 - Data_v8"
    doc.core_properties.subject = "多光谱无人机地块实例分割最终技术决策简报"
    doc.core_properties.author = "Project team"
    bullet_id = create_numbering(doc, ordered=False)
    number_id = create_numbering(doc, ordered=True)

    # memo_masthead pattern; report_title is a named 25 pt / INK override.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("最终模型技术决策简报")
    set_run_font(run, size=25, bold=True, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("多光谱无人机地块实例分割 · Data_v8 三 Seed 配对确认")
    set_run_font(run, size=13, color=MUTED)
    for label, value in (
        ("更新时间", "2026-08-17（Asia/Shanghai）"),
        ("评估范围", "train / valid；冻结 test 未读取"),
        ("当前状态", "保留部署模型；拒绝 C1 原始波段替换"),
        ("部署保护", f"SHA256 {summary['deployment_sha256']}；权重未修改"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        label_run = p.add_run(f"{label}：")
        set_run_font(label_run, size=10.5, bold=True, color=INK)
        value_run = p.add_run(value)
        set_run_font(value_run, size=10.5)
    add_callout(
        doc,
        "最终结论",
        "C1（RGB+NIR+RedEdge）三 seed 平均变化为 -0.006 pp，Seed 1 退化 1.336 pp；未通过平均提升与最差 seed 两项门控。正式部署继续使用 YOLO11n-seg + RGB+NDVI+NDRE + 3×CoordAtt + 3×ECA。",
    )

    doc.add_heading("1. 执行摘要", level=1)
    add_body_paragraph(
        doc,
        "Data_v8 阶段完成了 YOLOv8 低成本结构筛选、YOLO11 三 seed 配对确认、valid 逐图诊断和部署哈希复核。六次正式训练共记录 356 个 epoch，所有研究输出写入独立目录，部署 checkpoint 未被修改。",
    )
    for item in (
        "当前正式模型不是 YOLO11 与 YOLOv8 的推理混合；YOLOv8 仅用于结构/特征筛选。",
        "当前正式输入已经包含植被指数 NDVI 与 NDRE；原始 NIR/RedEdge 替换没有稳定收益。",
        "土壤代理 RE_BSI_PROXY、MCARI 以及二者组合均未通过 Seed-0 筛选，不进入部署。",
        "坡度尚未训练：缺少与影像严格对齐的 DEM/DSM/DTM，因此只能写“不可评估”。",
    ):
        add_list_item(doc, item, bullet_id)
    add_table(
        doc,
        ["关键项", "最终数值", "决策"],
        [
            ["当前部署参考 valid", "0.36905", "继续部署"],
            ["部署 valid 逐图诊断", "0.36648", "conf=0.25；用于效果图"],
            ["C1 三 seed 平均变化", "-0.006 pp", "未达 +0.300 pp 门槛"],
            ["C1 最差 seed", "-1.336 pp", "低于 -0.300 pp 底线"],
            ["部署权重", "SHA256 未变", "禁止自动写入"],
        ],
        [2520, 2520, 4320],
    )

    doc.add_page_break()
    doc.add_heading("2. 当前模型、输入特征与机制", level=1)
    add_body_paragraph(
        doc,
        "正式部署模型为 YOLO11n-seg，参数规模约 2.85M。骨干使用 Conv/C3k2、SPPF 与 C2PSA 形成多尺度语义表示，分割头在 P3/P4/P5 尺度联合预测实例框、类别与掩膜原型。",
    )
    add_table(
        doc,
        ["层面", "当前实现", "作用"],
        [
            ["输入", "R、G、B、NDVI、NDRE", "纹理/边界 + 植被活力 + 红边状态"],
            ["主干", "YOLO11n-seg", "多尺度实例检测与分割"],
            ["空间注意力", "3×CoordAtt", "加强方向、位置与长条作物行关系"],
            ["通道注意力", "3×ECA", "轻量重标定 RGB/指数通道贡献"],
            ["标签机制", "教师精炼 + 人工实例约束", "只修正匹配实例，不把教师幻觉直接写回"],
            ["稳健性", "Seed 0/1/2 配对", "区分稳定改进与偶然峰值"],
            ["部署机制", "冻结 test + SHA256 门控", "防止数据泄漏和误覆盖"],
        ],
        [1800, 3000, 4560],
    )
    add_callout(
        doc,
        "架构澄清",
        "YOLO8 和 YOLO11 没有在上线推理时混合。YOLO8 的作用是快速筛选；最终保留的部署推理图只包含 YOLO11n-seg 及其 CoordAtt/ECA 注意力。",
        fill="E9F3EE",
        color=GREEN,
    )

    doc.add_heading("3. 土壤、植被与坡度特征结论", level=1)
    add_table(
        doc,
        ["特征/机制", "实验结果", "最终处理"],
        [
            ["NDVI + NDRE", "当前部署持续使用", "保留；已经包含植被提取信息"],
            ["RE_BSI_PROXY", "Seed-0 较五通道参考 -5.832 pp", "未晋级；不是标准 SWIR-BSI"],
            ["MCARI", "Seed-0 较五通道参考 -5.723 pp", "未晋级"],
            ["RE_BSI_PROXY + MCARI", "Seed-0 较参考 -5.923 pp", "未证明互补"],
            ["NIR + RedEdge 替换", "YOLO11 三 seed 平均 -0.006 pp", "拒绝替换 NDVI + NDRE"],
            ["P4 双分支", "比 C1 仅 +0.004 pp", "未证明融合收益"],
            ["多尺度门控", "比 C1 -0.211 pp", "停止；F3 不运行"],
            ["坡度", "无严格对齐 DEM/DSM/DTM", "未训练、不可评价"],
        ],
        [2160, 3240, 3960],
    )
    add_body_paragraph(
        doc,
        "P4 Multispectral 不提供 SWIR 波段，因此项目中的 RE_BSI_PROXY 只能称为实验性红边裸土代理，不能宣称为标准 BSI 或 NDTI。筛选结果说明“可以计算”不等于“值得加入模型”。",
    )

    doc.add_page_break()
    doc.add_heading("4. YOLO11 三 Seed 配对确认", level=1)
    pair_rows = []
    for row in summary["paired_results"]:
        pair_rows.append(
            [
                str(row["seed"]),
                f"{float(row['c0_mask_map50_95']):.5f}",
                f"{float(row['c1_mask_map50_95']):.5f}",
                f"{float(row['delta_mask_map50_95']) * 100:+.3f} pp",
            ]
        )
    add_table(
        doc,
        ["Seed", "C0：RGB+NDVI+NDRE", "C1：RGB+NIR+RedEdge", "C1−C0"],
        pair_rows,
        [1080, 2880, 2880, 2520],
    )
    add_table(
        doc,
        ["门控条件", "阈值", "结果"],
        [
            ["三组配对完整", "3/3", "通过"],
            ["平均 Mask mAP50-95 变化", "≥ +0.300 pp", "-0.006 pp；未通过"],
            ["正向 seed 数", "≥ 2/3", "2/3；通过"],
            ["最差 seed", "≥ -0.300 pp", "-1.336 pp；未通过"],
            ["平均 Precision 回退", "≤ 1 pp", "-0.643 pp；通过"],
            ["平均 Recall 回退", "≤ 1 pp", "-0.041 pp；通过"],
        ],
        [3600, 2160, 3600],
    )
    add_callout(
        doc,
        "判定",
        "C1 在 Seed 0 提升 1.179 pp、Seed 2 提升 0.140 pp，但 Seed 1 退化 1.336 pp，导致均值接近零且方差较大。它表现出场景选择性，而不是可部署的稳定增益。",
    )

    doc.add_heading("5. 筛选机制为何停止", level=1)
    for item in (
        "YOLOv8 Seed-0 筛选中，C1 比 C0 提高 0.318 pp，因此有资格进入更严格的 YOLO11 三 seed 确认。",
        "P4 双分支 F1 仅比 C1 高 0.004 pp，收益小于预设门槛；多尺度门控 F2 比 C1 低 0.211 pp。",
        "由于 F2 未通过数值和困难切片门控，F3 按计划停止，避免继续消耗算力并扩大结构复杂度。",
        "最终 YOLO11 配对结果否定了 C1 的稳定替换价值，因此不再把原始波段或双分支写入部署模型。",
    ):
        add_list_item(doc, item, number_id)

    doc.add_page_break()
    doc.add_heading("6. 当前部署模型最新 valid 标注效果", level=1)
    add_body_paragraph(
        doc,
        "以下效果图来自当前正式部署权重在 60 张 valid 图上的重新诊断，输入为 RGB+NDVI+NDRE，imgsz 1024，confidence 0.25。图中同时保留强结果和典型失败，不读取冻结 test。",
    )
    add_picture(
        doc,
        DEPLOYMENT_IMAGE,
        "图 1  当前部署模型：人工标注（左）与模型预测（右）。上排为强结果，下排展示混合背景和低覆盖失败。",
        "Four valid samples comparing ground truth and the current deployed YOLO11 prediction, including two strong and two difficult cases.",
    )
    deploy_rows = []
    for row in summary["deployment_cases"]:
        deploy_rows.append(
            [
                str(row["image"]).replace(".JPG", ""),
                str(row["diagnosis"]),
                f"F1 {float(row['f1']):.3f}",
                f"FP {int(row['fp'])} / FN {int(row['fn'])}",
            ]
        )
    add_table(
        doc,
        ["样本", "诊断", "逐图结果", "错误"],
        deploy_rows,
        [1800, 2520, 2160, 2880],
    )
    add_body_paragraph(
        doc,
        "效果较好的 DJI_2350、DJI_1790 能连续覆盖主要作物行；DJI_1500 在屋顶/设施边缘产生大面积误检，DJI_2260 在极低覆盖场景召回不足，DJI_2330 则需要先人工确认空标注是否真实。",
    )

    doc.add_page_break()
    doc.add_heading("7. C1 方案：效果好与效果差的位置", level=1)
    add_body_paragraph(
        doc,
        "Seed 1 是三 seed 门控中的决定性失败 seed。下图使用同一人工真值，比较 C0 指数输入与 C1 原始波段输入。上排展示最大退化，下面展示一个真实改善，说明 C1 并非处处无效，而是缺乏整体稳定性。",
    )
    add_picture(
        doc,
        C1_IMAGE,
        "图 2  C1 代表性退化与改善：每行依次为人工标注、C0（RGB+NDVI+NDRE）和 C1（RGB+NIR+RedEdge）。",
        "Two valid samples comparing ground truth, C0 vegetation-index input, and C1 raw-band input; one regression and one improvement.",
    )
    compare_rows = []
    for row in summary["comparison_cases"]:
        compare_rows.append(
            [
                str(row["image"]).replace(".JPG", ""),
                str(row["diagnosis"]),
                f"{float(row['c0_f1']):.3f} → {float(row['c1_f1']):.3f}",
                f"{float(row['delta']):+.3f}",
                f"FP {int(row['fp_delta']):+d} / FN {int(row['fn_delta']):+d}",
            ]
        )
    add_table(
        doc,
        ["样本", "位置", "C0→C1 F1", "变化", "错误变化"],
        compare_rows,
        [1440, 2160, 2160, 1440, 2160],
    )
    add_callout(
        doc,
        "失败模式",
        "C1 的优势主要出现在个别低覆盖或设施干扰图，可能来自原始 NIR/RedEdge 对背景反射差异的直接响应；但在 DJI_3230、DJI_0030、DJI_1790 等密集作物行中，FP/FN 同时增加。该不均衡收益不足以替换更稳定的 NDVI/NDRE。",
        fill="FDECEC",
        color=RED,
    )

    doc.add_page_break()
    doc.add_heading("8. 下一轮工作计划", level=1)
    add_table(
        doc,
        ["优先级", "动作", "完成标准"],
        [
            ["P0", "复核 DJI_3230、DJI_0030、DJI_1790 的密集边界与标签", "区分真实漏检、相邻实例粘连和标签缺口"],
            ["P0", "建立道路/设施边缘困难负样本子集", "valid 逐图 FP/FN 与场景标签可复现"],
            ["P1", "保留五通道与注意力，测试边界损失或困难样本采样", "均值、最差 seed 和困难切片同时改善"],
            ["P1", "审计空标注 DJI_2330 等样本", "确认空标注真实性后再判断误检"],
            ["停止项", "不追加未晋级土壤/植被指数，不运行 F3", "避免无证据增维和结构膨胀"],
            ["待条件", "取得严格对齐 DEM/DSM/DTM 后再设计坡度实验", "空间对齐、分辨率与坐标系统审计通过"],
        ],
        [1080, 3780, 4500],
    )
    doc.add_heading("9. 审计与复现边界", level=1)
    for item in (
        "本轮正式确认使用 train/valid；test 未枚举、未读取、未用于调参。",
        "六次 YOLO11 训练使用 seed 0/1/2、60 epochs、patience 20、imgsz 1024、batch 2；共记录 356 个 epoch。",
        "每个 seed 的 C0/C1 从同一 yolo11n-seg.pt 和相同显式随机种子初始化，比较差异只来自第 4/5 通道内容。",
        f"部署 checkpoint SHA256 前后均为 {summary['deployment_sha256']}。",
        "研究 checkpoint 仅作为证据保留；C1 门控失败后不生成候选部署权重。",
    ):
        add_list_item(doc, item, bullet_id)
    add_callout(
        doc,
        "交付状态",
        "当前正式模型、特征和部署权重均不变。网页与本 Word 文档已把效果不佳的位置、未晋级特征以及未评估坡度分别标明。",
        fill="E9F3EE",
        color=GREEN,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"docx={OUTPUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
