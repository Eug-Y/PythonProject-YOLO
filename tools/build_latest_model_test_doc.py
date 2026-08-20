"""Build a visually verified Word summary from frozen model-test artifacts."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
REPORT_ROOT = ROOT / "reports" / "Latest_model_test_presentation"
SOURCE = ROOT / "reports" / "Data_v3_spectral_aug_test_after_error_review"
ASSETS = REPORT_ROOT / "doc_assets"
OUTPUT = REPORT_ROOT / "最新模型测试效果总结.docx"

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"
BLACK = RGBColor(0x17, 0x20, 0x2A)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
GREEN = RGBColor(0x16, 0x5F, 0x39)
GRAY = RGBColor(0x63, 0x6E, 0x67)
LIGHT_GRAY = "F2F4F7"
LIGHT_GREEN = "ECF7F0"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_run_font(run, *, size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_width(parent, tag: str, width_dxa: int) -> None:
    width = ensure_child(parent, tag)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(width_dxa))


def apply_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError(f"Table widths must total {TABLE_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    set_width(tbl_pr, "w:tblW", TABLE_WIDTH_DXA)
    indent = ensure_child(tbl_pr, "w:tblInd")
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for col_idx, width in enumerate(widths):
        table.columns[col_idx].width = Twips(width)
    for row in table.rows:
        row.height = None
        for col_idx, cell in enumerate(row.cells):
            width = widths[col_idx]
            cell.width = Twips(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            set_width(tc_pr, "w:tcW", width)
            tc_mar = ensure_child(tc_pr, "w:tcMar")
            for side, value in CELL_MARGINS.items():
                margin = ensure_child(tc_mar, f"w:{side}")
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = FONT_LATIN
    caption.font.size = Pt(9.5)
    caption.font.italic = False
    caption.font.color.rgb = GRAY
    caption._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def set_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("模型测试效果总结 | 第 ")
    set_run_font(run, size=9, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=GRAY)


def configure_section(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("YOLO 多光谱实例分割 | 独立测试总结")
    set_run_font(run, size=9, color=GRAY, bold=True)
    set_page_number(section.footer.paragraphs[0])


def add_metadata(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.10
    label_run = p.add_run(f"{label}：")
    set_run_font(label_run, size=10.5, color=BLACK, bold=True)
    value_run = p.add_run(value)
    set_run_font(value_run, size=10.5, color=BLACK)


def add_callout(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.line_spacing = 1.15
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GREEN)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "7")
    left.set(qn("w:color"), "2B8A57")
    borders.append(left)
    p_pr.append(borders)
    lead = p.add_run(f"{label}  ")
    set_run_font(lead, size=11, color=GREEN, bold=True)
    body = p.add_run(text)
    set_run_font(body, size=11, color=BLACK)


def style_table_text(table) -> None:
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.10
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, size=8.6, color=BLACK, bold=(row_idx == 0))
            if row_idx == 0:
                set_cell_fill(cell, LIGHT_GRAY)
            elif row_idx == 1:
                set_cell_fill(cell, LIGHT_GREEN)


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_model_table(doc: Document) -> None:
    columns = ["模型", "定位", "Mask P", "Mask R", "Mask mAP50", "Mask mAP50-95"]
    rows = [
        ["spectral_aug_seed0_v2", "单项最高", "0.6843", "0.5842", "0.5800", "0.3915"],
        ["grouped_reviewed_v1", "当前部署", "0.7000", "0.5776", "0.5810", "0.3886"],
        ["v5 RGB+MCARI", "最新研究", "0.6865", "0.5568", "0.5647", "0.3782"],
        ["v4 RGB+RE_BSI", "研究归档", "0.6716", "0.5794", "0.5730", "0.3774"],
    ]
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for idx, value in enumerate(columns):
        table.rows[0].cells[idx].text = value
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    apply_table_geometry(table, [2500, 1250, 1050, 1050, 1550, 1960])
    mark_header_row(table.rows[0])
    style_table_text(table)
    p = doc.add_paragraph("表 1  冻结独立测试结果对比；主要排序指标为 Mask mAP50-95。", style="Caption")
    p.paragraph_format.keep_with_next = False


def pil_font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc") if bold else Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
    ]
    for path in candidates:
        if path.is_file():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def make_comparison(batch: int) -> Path:
    left = Image.open(SOURCE / f"val_batch{batch}_labels.jpg").convert("RGB")
    right = Image.open(SOURCE / f"val_batch{batch}_pred.jpg").convert("RGB")
    panel_w = 1024
    panel_h = 1024
    left = ImageOps.fit(left, (panel_w, panel_h), method=Image.Resampling.LANCZOS)
    right = ImageOps.fit(right, (panel_w, panel_h), method=Image.Resampling.LANCZOS)
    top_h = 78
    canvas = Image.new("RGB", (panel_w * 2 + 6, panel_h + top_h), "white")
    draw = ImageDraw.Draw(canvas)
    draw.rectangle((0, 0, panel_w, top_h), fill="#EAF2F8")
    draw.rectangle((panel_w + 6, 0, panel_w * 2 + 6, top_h), fill="#EAF7EF")
    draw.text((34, 19), "人工标注（Ground Truth）", font=pil_font(34, True), fill="#24527A")
    draw.text((panel_w + 40, 19), "模型预测（Prediction）", font=pil_font(34, True), fill="#1D6B43")
    canvas.paste(left, (0, top_h))
    canvas.paste(right, (panel_w + 6, top_h))
    draw.rectangle((0, top_h, panel_w - 1, panel_h + top_h - 1), outline="#9BB8CE", width=4)
    draw.rectangle((panel_w + 6, top_h, panel_w * 2 + 5, panel_h + top_h - 1), outline="#90C5A8", width=4)
    output = ASSETS / f"test_batch_{batch + 1}_comparison.jpg"
    canvas.save(output, quality=94, subsampling=0)
    return output


def make_diagnostics() -> Path:
    inputs = [
        (SOURCE / "MaskPR_curve.png", "Mask PR 曲线"),
        (SOURCE / "MaskF1_curve.png", "Mask F1 曲线"),
        (SOURCE / "confusion_matrix_normalized.png", "归一化混淆矩阵"),
    ]
    panels = []
    panel_h = 780
    for path, label in inputs:
        image = Image.open(path).convert("RGB")
        ratio = panel_h / image.height
        resized = image.resize((round(image.width * ratio), panel_h), Image.Resampling.LANCZOS)
        panels.append((resized, label))
    header_h = 72
    gap = 16
    total_w = sum(image.width for image, _ in panels) + gap * (len(panels) - 1)
    canvas = Image.new("RGB", (total_w, panel_h + header_h), "white")
    draw = ImageDraw.Draw(canvas)
    x = 0
    for image, label in panels:
        draw.rectangle((x, 0, x + image.width, header_h), fill="#EEF3F0")
        bbox = draw.textbbox((0, 0), label, font=pil_font(28, True))
        text_w = bbox[2] - bbox[0]
        draw.text((x + (image.width - text_w) / 2, 18), label, font=pil_font(28, True), fill="#2C5E45")
        canvas.paste(image, (x, header_h))
        draw.rectangle((x, header_h, x + image.width - 1, panel_h + header_h - 1), outline="#CAD6CF", width=3)
        x += image.width + gap
    output = ASSETS / "test_diagnostics.jpg"
    canvas.save(output, quality=94, subsampling=0)
    return output


def set_picture_alt(paragraph, description: str) -> None:
    doc_pr = paragraph._p.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[-1].set("descr", description)


def add_figure(doc: Document, image_path: Path, caption: str, alt: str, width: float = 6.35) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(image_path), width=Inches(width))
    set_picture_alt(p, alt)
    cap = doc.add_paragraph(caption, style="Caption")
    cap.paragraph_format.keep_with_next = False


def build_document() -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    comparisons = [make_comparison(i) for i in range(3)]
    diagnostics = make_diagnostics()

    doc = Document()
    configure_styles(doc)
    configure_section(doc)
    props = doc.core_properties
    props.title = "最新模型测试效果总结"
    props.subject = "YOLO 多光谱实例分割模型独立测试结果与效果图"
    props.author = "项目团队"

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("最新模型测试效果总结")
    set_run_font(run, size=23, color=BLACK, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("YOLO 多光谱实例分割 · 冻结独立测试结果")
    set_run_font(run, size=14, color=GRAY)

    add_metadata(doc, "测试集", "Data_v3 test；30 张图片，1681 个实例")
    add_metadata(doc, "测试设置", "1024 px，batch 1；结果来自已保存的冻结评估")
    add_metadata(doc, "主要指标", "Mask mAP50-95")
    add_metadata(doc, "整理日期", "2026-08-08")

    add_callout(
        doc,
        "核心结论",
        "单次测试中 spectral_aug_seed0_v2 的 Mask mAP50-95 最高（0.3915）；但仅比当前部署模型高 0.2854 个百分点，且当前模型的 Mask Precision 与 Mask mAP50 更高。因此将其作为单项指标冠军展示，部署模型暂不切换。最新 MCARI 候选为 0.3782，未达到升级条件。",
    )

    doc.add_heading("1. 冻结测试结果", level=1)
    add_model_table(doc)
    p = doc.add_paragraph()
    run = p.add_run("最高单项模型：")
    set_run_font(run, bold=True, color=GREEN)
    run = p.add_run("exp_ca_eca_5ch_n_spectral_aug_seed0_v2，输入 R、G、B、NDVI、NDRE。")
    set_run_font(run)
    p = doc.add_paragraph()
    run = p.add_run("部署建议：")
    set_run_font(run, bold=True, color=BLUE)
    run = p.add_run("继续使用 exp_ca_5ch_grouped_reviewed_v1；待多随机种子结果证明提升稳定后再考虑切换。")
    set_run_font(run)

    doc.add_page_break()
    doc.add_heading("2. 测试集预测效果", level=1)
    p = doc.add_paragraph("以下对照图直接复用冻结测试时保存的输出。左侧是人工标注，右侧是模型预测；没有重新读取 test 调参，也没有重复计算最终指标。")
    p.paragraph_format.keep_with_next = True
    add_figure(doc, comparisons[0], "图 1  测试批次 1：主体作物行基本识别，密集相邻轮廓仍有合并或漏检。", "测试批次1人工标注与模型预测对照")
    add_figure(doc, comparisons[1], "图 2  测试批次 2：主要边界整体对齐，局部小轮廓和相邻实例仍有偏差。", "测试批次2人工标注与模型预测对照")

    doc.add_page_break()
    doc.add_heading("3. 更多效果与诊断", level=1)
    add_figure(doc, comparisons[2], "图 3  测试批次 3：上部作物行识别较完整，但下部道路/背景出现明显误检。", "测试批次3人工标注与模型预测对照")
    add_figure(doc, diagnostics, "图 4  Mask PR、Mask F1 与归一化混淆矩阵。", "最高单项模型的Mask PR曲线、Mask F1曲线和归一化混淆矩阵", width=6.2)

    doc.add_heading("4. 结论与下一步", level=1)
    doc.add_paragraph("从测试图可以看到，模型能够提取大部分连续地块轮廓，但密集小目标、相互贴近的轮廓以及边界模糊区域仍存在漏检、粘连或边界偏移；道路、裸地等非目标背景也可能触发较大面积误检。继续堆叠相似光谱指数没有带来稳定泛化收益。")
    doc.add_paragraph("下一阶段应优先进行多随机种子复验、困难样本定向补充和边界标注一致性检查；在这些证据完成前，不建议仅凭单次测试的 0.2854 个百分点差异替换部署模型。")

    doc.add_heading("数据来源", level=2)
    source = doc.add_paragraph()
    source.paragraph_format.space_before = Pt(4)
    source.paragraph_format.space_after = Pt(4)
    source.paragraph_format.line_spacing = 1.10
    run = source.add_run(
        "reports/Data_v3_spectral_aug_comparison/summary.json；"
        "reports/Data_v4_final_model_comparison/summary.json；"
        "reports/Data_v5_final_model_comparison/summary.json；"
        "reports/Data_v3_spectral_aug_test_after_error_review/。"
    )
    set_run_font(run, size=9, color=GRAY)

    REPORT_ROOT.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
