from __future__ import annotations

import argparse
import importlib.util
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def load_geometry_helper(path: Path):
    spec = importlib.util.spec_from_file_location("table_geometry", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load helper: {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def mark_header_row(table) -> None:
    row = table.rows[0]
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        tr_pr.append(marker)
    marker.set(qn("w:val"), "true")


def grid_widths(table) -> list[int]:
    widths = []
    for column in table._tbl.tblGrid:
        value = column.get(qn("w:w"))
        if value:
            widths.append(int(value))
    if len(widths) != len(table.columns):
        raise RuntimeError("Could not resolve complete table grid")
    return widths


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--helper", type=Path, required=True)
    args = parser.parse_args()

    geometry = load_geometry_helper(args.helper)
    document = Document(args.docx)
    for table in document.tables:
        mark_header_row(table)
        widths = grid_widths(table)
        if sum(widths) > 9360:
            widths = geometry.column_widths_from_weights(widths, 9360)
        geometry.apply_table_geometry(
            table,
            widths,
            indent_dxa=100,
            cell_margins_dxa={"top": 80, "bottom": 80, "start": 100, "end": 100},
        )
    document.save(args.docx)
    print(f"FINALIZED_TABLES={len(document.tables)}")


if __name__ == "__main__":
    main()
