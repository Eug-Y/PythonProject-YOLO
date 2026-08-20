"""Build the GitHub multispectral segmentation reference guide DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "reports" / "GitHub_多光谱与航拍分割项目参考指南.docx"

LATIN_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "243447"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
BORDER = "B8C4D1"
GREEN = "317A4B"
GOLD = "8A6400"


PROJECTS = [
    {
        "name": "YOLOv11-RGBT",
        "url": "https://github.com/wandahangFY/YOLOv11-RGBT",
        "summary": "YOLO11 多光谱实例分割框架，支持任意通道，以及早期、中期、后期、分数和权重共享等融合方式。适合研究 RGB 与 NDVI/NDRE 双分支融合。",
        "priority": "高",
        "license": "AGPL-3.0",
    },
    {
        "name": "YOLO-MIF",
        "url": "https://github.com/wandahangFY/YOLO-MIF",
        "summary": "包含多通道检测、分割和多种信息融合策略，许可证宽松，便于借鉴数据加载和融合模块；主体基于较旧的 YOLOv8。",
        "priority": "中高",
        "license": "MIT",
    },
    {
        "name": "Ultralytics 多光谱分割讨论",
        "url": "https://github.com/ultralytics/ultralytics/issues/22620",
        "summary": "官方确认多光谱 TIFF 可用于分割任务；数据 YAML 中需要正确配置 channels，并按官方多光谱数据格式组织输入。",
        "priority": "高",
        "license": "官方问题记录",
    },
    {
        "name": "SAHI",
        "url": "https://github.com/obss/sahi",
        "summary": "面向大图的小目标切片推理框架，支持目标检测和实例分割，并可接入 Ultralytics。适合一张航拍图包含大量小轮廓的场景。",
        "priority": "高",
        "license": "MIT",
    },
    {
        "name": "FiftyOne",
        "url": "https://github.com/voxel51/fiftyone",
        "summary": "用于可视化评估、错误分析、嵌入检索和数据集质量管理，可自动定位误检、漏检与困难样本，支撑主动学习闭环。",
        "priority": "中高",
        "license": "Apache-2.0",
    },
    {
        "name": "Kornia",
        "url": "https://github.com/kornia/kornia",
        "summary": "基于 PyTorch 的几何视觉与增强库，可实现任意通道图像与分割轮廓同步变换，适合编写 RGB、NDVI、NDRE 分通道增强。",
        "priority": "高",
        "license": "Apache-2.0",
    },
    {
        "name": "Satellite Image Deep Learning",
        "url": "https://github.com/satellite-image-deep-learning/techniques",
        "summary": "航拍与卫星影像深度学习技术合集，覆盖遥感分割、目标检测、切片、嵌入和多光谱方法，适合持续查找方案与基准。",
        "priority": "参考",
        "license": "Apache-2.0",
    },
]


def set_run_font(run, *, size=None, color=None, bold=None, italic=None, underline=None):
    run.font.name = LATIN_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(index, len(widths) - 1)])
            set_cell_margins(cell)


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_hyperlink(paragraph, text, url, *, bold=False, size=None):
    relationship_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    run_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.append(underline)
    if bold:
        bold_node = OxmlElement("w:b")
        run_pr.append(bold_node)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), CJK_FONT)
    run_pr.append(fonts)
    if size is not None:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        run_pr.append(sz)
        sz_cs = OxmlElement("w:szCs")
        sz_cs.set(qn("w:val"), str(int(size * 2)))
        run_pr.append(sz_cs)
    run.append(run_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_paragraph_border_and_fill(paragraph, *, fill=CALLOUT, border_color=BLUE):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border_color)
    borders.append(left)
    p_pr.append(borders)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "180")
    ind.set(qn("w:right"), "120")
    p_pr.append(ind)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Title", 28, INK, 0, 8),
        ("Subtitle", 14, MUTED, 0, 18),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = LATIN_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if style_name != "Subtitle":
            style.font.bold = style_name != "Title" or True


def add_numbering_definition(doc, num_type):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if num_type == "bullet" else "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if num_type == "bullet" else "%1.")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), LATIN_FONT)
    r_fonts.set(qn("w:hAnsi"), LATIN_FONT)
    r_fonts.set(qn("w:eastAsia"), CJK_FONT)
    r_pr.append(r_fonts)
    level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, *, bold_prefix=None):
    paragraph = doc.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        prefix_run = paragraph.add_run(bold_prefix)
        set_run_font(prefix_run, bold=True)
        body_run = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(body_run)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def set_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=9, color=MUTED)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("技术参考指南  |  多光谱与航拍实例分割")
    set_run_font(hr, size=8.5, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    set_page_field(fp)


def add_project_detail(doc, project, bullet_num_id, use_text, caution_text):
    heading = doc.add_paragraph(style="Heading 2")
    add_hyperlink(heading, project["name"], project["url"], bold=True, size=13)
    meta = heading.add_run(f"  ·  {project['license']}")
    set_run_font(meta, size=9.5, color=MUTED, bold=False)
    add_list_item(doc, f"适用价值：{project['summary']}", bullet_num_id, bold_prefix="适用价值：")
    add_list_item(doc, f"落地建议：{use_text}", bullet_num_id, bold_prefix="落地建议：")
    add_list_item(doc, f"注意事项：{caution_text}", bullet_num_id, bold_prefix="注意事项：")


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    bullet_num_id = add_numbering_definition(doc, "bullet")
    decimal_num_id = add_numbering_definition(doc, "decimal")

    props = doc.core_properties
    props.title = "GitHub 多光谱与航拍分割项目参考指南"
    props.subject = "面向 YOLO11 五通道农田实例分割的开源项目调研"
    props.author = ""
    props.keywords = "YOLO11, multispectral, segmentation, UAV, NDVI, NDRE"

    # Editorial cover pattern.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(48)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    run = kicker.add_run("开源技术调研 · 2026")
    set_run_font(run, size=10.5, color=GOLD, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("GitHub 多光谱与航拍分割\n项目参考指南")
    set_run_font(run, size=28, color=INK, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("面向 YOLO11 五通道农田实例分割\nRGB + NDVI + NDRE")
    set_run_font(run, size=14, color=MUTED)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(10)
    meta.paragraph_format.space_after = Pt(30)
    run = meta.add_run("整理日期：2026年7月31日  |  项目数量：7  |  用途：架构选型与实验规划")
    set_run_font(run, size=9.5, color=MUTED)

    lead = doc.add_paragraph()
    lead.paragraph_format.space_before = Pt(8)
    lead.paragraph_format.space_after = Pt(10)
    lead.paragraph_format.line_spacing = 1.25
    add_paragraph_border_and_fill(lead)
    label = lead.add_run("核心建议  ")
    set_run_font(label, bold=True, color=DARK_BLUE)
    text = lead.add_run(
        "优先验证 Ultralytics 原生多光谱数据链路，随后对比五通道早期融合与 RGB/指数双分支中期融合；"
        "切片推理、困难样本挖掘和通道安全增强用于进一步提升小轮廓召回率与泛化能力。"
    )
    set_run_font(text)

    source_note = doc.add_paragraph()
    source_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source_note.paragraph_format.space_before = Pt(12)
    run = source_note.add_run("资料来源：GitHub 仓库与 Ultralytics 官方问题记录")
    set_run_font(run, size=9, color=MUTED, italic=True)
    doc.add_page_break()

    # Comparison table.
    doc.add_heading("1. 项目对比速览", level=1)
    intro = doc.add_paragraph("以下项目按与当前五通道农田实例分割任务的直接相关性整理。优先级表示建议首先验证的顺序，不代表仓库质量排名。")
    intro.paragraph_format.space_after = Pt(8)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [2160, 4536, 1296, 1368]
    headers = ["项目", "可借鉴内容", "优先级", "许可证/来源"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run_font(run, size=9.5, bold=True, color=DARK_BLUE)
    set_repeat_table_header(table.rows[0])

    for project in PROJECTS:
        row = table.add_row()
        prevent_row_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        name_p = row.cells[0].paragraphs[0]
        name_p.paragraph_format.space_after = Pt(0)
        add_hyperlink(name_p, project["name"], project["url"], bold=True, size=9.5)
        summary_p = row.cells[1].paragraphs[0]
        summary_p.paragraph_format.space_after = Pt(0)
        summary_p.paragraph_format.line_spacing = 1.12
        run = summary_p.add_run(project["summary"])
        set_run_font(run, size=9.2)
        priority_p = row.cells[2].paragraphs[0]
        priority_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        priority_p.paragraph_format.space_after = Pt(0)
        run = priority_p.add_run(project["priority"])
        priority_color = GREEN if project["priority"] == "高" else GOLD if "中" in project["priority"] else MUTED
        set_run_font(run, size=9.5, bold=True, color=priority_color)
        license_p = row.cells[3].paragraphs[0]
        license_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        license_p.paragraph_format.space_after = Pt(0)
        run = license_p.add_run(project["license"])
        set_run_font(run, size=8.7, color=MUTED)

    set_table_geometry(table, widths)
    set_table_borders(table)
    caption = doc.add_paragraph("表1  GitHub 多光谱与航拍分割项目对比")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)
    run = caption.runs[0]
    set_run_font(run, size=9, color=MUTED, italic=True)

    doc.add_heading("2. 架构与多光谱输入", level=1)
    add_project_detail(
        doc,
        PROJECTS[0],
        bullet_num_id,
        "将 RGB 作为主分支、NDVI/NDRE 作为指数分支，在骨干网络中段融合；与当前五通道早期融合做严格 A/B 对照。",
        "仓库为 AGPL-3.0。直接复制到闭源或商业系统前需要评估许可证；建议优先借鉴结构思想与实验设计。",
    )
    add_project_detail(
        doc,
        PROJECTS[1],
        bullet_num_id,
        "重点阅读多模态数据加载、通道配置和中间融合模块；可作为实现简化版双分支模型的参考。",
        "主体基于 YOLOv8，与当前 YOLO11 代码结构不同，不建议整仓替换现有工程。",
    )
    add_project_detail(
        doc,
        PROJECTS[2],
        bullet_num_id,
        "验证当前 Ultralytics 版本是否可以直接使用 channels: 5 与官方多光谱 TIFF 数据格式，减少自定义加载器维护成本。",
        "官方示例偏向多页 TIFF。当前 NPY 数据需要继续使用适配器，或先做小规模格式转换验证。",
    )

    doc.add_heading("3. 航拍推理、增强与数据闭环", level=1)
    add_project_detail(
        doc,
        PROJECTS[3],
        bullet_num_id,
        "先在验证集测试 768-1024 像素切片、20%-25% 重叠，并对跨切片实例进行合并；重点观察小轮廓漏检和边界完整性。",
        "当前模型是自定义五通道输入，需要编写适配层；大面积跨切片目标可能被切碎，必须单独检查。",
    )
    add_project_detail(
        doc,
        PROJECTS[4],
        bullet_num_id,
        "导入每张图的预测、置信度和标签，按误检、漏检、低置信度和相似度聚类生成下一轮复核队列。",
        "它主要提升数据分析效率，不会直接提高模型分数；需要先定义清晰的错误分类与复核规则。",
    )
    add_project_detail(
        doc,
        PROJECTS[5],
        bullet_num_id,
        "实现所有通道与轮廓同步的旋转、翻转、缩放；颜色扰动只作用于 RGB，NDVI/NDRE 使用轻微增益、偏移和噪声。",
        "不要把 RGB 的 HSV 变换直接作用于指数通道；所有光谱通道必须保持像素对齐。",
    )
    add_project_detail(
        doc,
        PROJECTS[6],
        bullet_num_id,
        "作为后续检索入口，重点关注遥感实例分割、图像切片、特征融合和数据集划分相关条目。",
        "这是资源合集而非可直接运行的单一框架，采用其中方案前仍需检查原始仓库和论文。",
    )

    doc.add_heading("4. 建议落地顺序", level=1)
    steps = [
        "原生多光谱链路验证：用少量样本验证 channels: 5、数据读取、训练、验证和推理结果一致。",
        "通道安全增强：补充上下翻转、90°旋转、RGB 专属颜色扰动和指数通道轻微噪声。",
        "融合架构对照：保持数据划分与超参数一致，对比五通道早期融合和 RGB/指数中期融合。",
        "模型容量对照：在显存允许时增加 YOLO11s-seg，对比当前 YOLO11n-seg。",
        "切片推理实验：使用 SAHI 思路验证小轮廓召回率和跨切片边界。",
        "困难样本闭环：使用 FiftyOne 思路持续筛选误检、漏检和代表性不足的飞行批次。",
    ]
    for step in steps:
        add_list_item(doc, step, decimal_num_id)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note.paragraph_format.space_after = Pt(6)
    add_paragraph_border_and_fill(note, fill="FFF8E8", border_color=GOLD)
    label = note.add_run("评估原则  ")
    set_run_font(label, bold=True, color=GOLD)
    body = note.add_run(
        "每组实验至少运行 3 个随机种子，并使用从未参与任何训练和调参的新飞行批次作为冻结测试集。"
        "仅当 Mask mAP50-95、精确率和召回率在干净测试集上稳定提升时，才替换网站默认模型。"
    )
    set_run_font(body)

    doc.add_heading("5. GitHub 来源", level=1)
    for project in PROJECTS:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(decimal_num_id))
        num_pr.append(ilvl)
        num_pr.append(num_id_node)
        p_pr.append(num_pr)
        add_hyperlink(p, project["name"], project["url"], bold=True)
        tail = p.add_run(f"  —  {project['url']}")
        set_run_font(tail, size=9, color=MUTED)

    closing = doc.add_paragraph()
    closing.paragraph_format.space_before = Pt(12)
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = closing.add_run("文档结束")
    set_run_font(run, size=9, color=MUTED)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
